import re
from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# GoodEarth Apparels Ltd. — Buyer: Joe Fresh
# 'Booking date ... Shipping mark as below' Excel ফরম্যাট।
#   - একটা ফাইলে একাধিক শিট থাকতে পারে, আর প্রতিটা শিটে একাধিক 'ব্লক' (প্রতিটা
#     ব্লক একটা কালার/PO-র বুকিং, ~১৯ রো) থাকে। প্রতিটা ব্লক 'Booking date :'
#     লেখা দিয়ে শুরু হয় — সেই লেখা দেখেই ব্লক আলাদা করা হয়, কাজেই ব্লক-সংখ্যা
#     যত-ই হোক ঠিকমতো কাজ করে। ব্লকের ভেতরের সব তথ্য লেবেল-টেক্সট দেখে খোঁজা
#     হয় (fixed রো-নম্বর ধরা হয় না)।
#   - ব্লক থেকে (প্রতিটা ব্লক = ১টা Master Carton লাইন):
#       'Style -S7WP080425'          -> Gmt. Style No
#       'Po No: 4885286927 32'       -> Gmt. PO (কোলনের পরের পুরো লেখা)
#       'Color name' (C1) / CODE/COLOR -> Reference/SKU Number
#       'UPC CODE : 2088...'         -> Pack Type
#       'Carton measuremnet:- L-15'' X W-14" X H-5" inch' -> Length/Width/
#                                       Height, একক ইঞ্চি (Measurement Unit
#                                       = 'Inch', কোনো কনভার্সন না)
#       'Need Cartoon Qty:' এর পাশের সংখ্যা -> Order Qty
#   - Item Name/Ply UI থেকে সিলেক্ট করা।
# ---------------------------------------------------------------------------


def _norm(s):
    return re.sub(r'[^a-z0-9]', '', str(s or '').lower())


def _clean(v):
    if v is None:
        return ''
    return re.sub(r'\s+', ' ', str(v)).strip()


def _is_num(v):
    if v is None:
        return False
    try:
        float(v)
        return True
    except (TypeError, ValueError):
        return False


def _fmt_num(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return ''
    return str(int(f)) if f == int(f) else str(round(f, 3))


# L-15'' X W-14“ X H-5” inch — কোট-চিহ্নগুলো নানা রকম (' '' “ ” ") হতে পারে,
# তাই সংখ্যার মাঝে যেকোনো non-digit অক্ষর সহ্য করা হয়।
_MEAS_RE = re.compile(
    r'L\s*[-:]?\s*(\d+(?:\.\d+)?)\D*?W\s*[-:]?\s*(\d+(?:\.\d+)?)\D*?H\s*[-:]?\s*(\d+(?:\.\d+)?)',
    re.I,
)


def _parse_measurement(text):
    """রিটার্ন করে (l, w, h, unit) — unit: 'Inch' / 'Cm' / None (বোঝা যায়নি)।
    mm হলে cm-এ কনভার্ট করে 'Cm' দেওয়া হয় (সিস্টেমের সাধারণ নিয়ম)।"""
    if not text:
        return None
    m = _MEAS_RE.search(str(text))
    if not m:
        return None
    l, w, h = float(m.group(1)), float(m.group(2)), float(m.group(3))
    t = str(text).lower()
    if 'inch' in t or re.search(r'\bin\b', t):
        unit = 'Inch'
    elif re.search(r'\bmm\b', t):
        l, w, h, unit = l / 10, w / 10, h / 10, 'Cm'
    elif 'cm' in t:
        unit = 'Cm'
    else:
        unit = None
    return _fmt_num(l), _fmt_num(w), _fmt_num(h), unit


def _value_after_colon(text):
    """'UPC CODE   :  2088...' -> '2088...'। কোলন না থাকলে ''।"""
    s = str(text)
    if ':' in s:
        return _clean(s.split(':', 1)[1])
    return ''


def _right_value(ws, r, c, numeric=False):
    """সেলের ডানদিকের প্রথম নন-এম্পটি সেল (numeric=True হলে প্রথম সংখ্যা)।"""
    for c2 in range(c + 1, ws.max_column + 1):
        v = ws.cell(row=r, column=c2).value
        if v is None or _clean(v) == '':
            continue
        if numeric and not _is_num(v):
            continue
        return v
    return None


def _find_blocks(ws, max_rows):
    """'Booking date' দিয়ে শুরু হওয়া রো-গুলো (কলাম A) থেকে ব্লকের (start,
    end) তালিকা। কোনো মার্কার না পেলে পুরো শিট একটাই ব্লক।"""
    starts = []
    blank_run = 0
    for r in range(1, max_rows + 1):
        v = ws.cell(row=r, column=1).value
        if v is None or _clean(v) == '':
            blank_run += 1
            if blank_run >= 200:
                max_rows = r
                break
            continue
        blank_run = 0
        if _norm(v).startswith('bookingdate'):
            starts.append(r)
    if not starts:
        return [(1, max_rows)]
    blocks = []
    for i, s in enumerate(starts):
        e = (starts[i + 1] - 1) if i + 1 < len(starts) else max_rows
        blocks.append((s, e))
    return blocks


def _read_block(ws, start, end, sheet_name, filename, item_name, ply_value):
    """একটা ব্লক থেকে (line_item_or_None, warnings)।"""
    warnings = []
    style = po = color = upc = ''
    meas = None
    qty_val = None
    color_fallback = ''
    meas_fallback = None

    for r in range(start, end + 1):
        for c in range(1, min(ws.max_column, 8) + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            n = _norm(v)
            if not n:
                continue
            raw = str(v)

            if c == 1 and n.startswith('style') and not style:
                style = _clean(re.sub(r'^\s*style\s*[-:]?\s*', '', raw, flags=re.I))
            elif c == 1 and n.startswith('pono') and not po:
                po = _clean(re.sub(r'^\s*po\s*no\s*[:.\-]?\s*', '', raw, flags=re.I))
            elif c == 1 and n.startswith('cartonmeas') and meas is None:
                meas = _parse_measurement(raw)
            elif n == 'colorname' and not color:
                rv = _right_value(ws, r, c)
                color = _clean(rv)
            elif n.startswith('codecolor') and not color_fallback:
                color_fallback = _value_after_colon(raw)
            elif n.startswith('ctnmeas') and meas_fallback is None:
                meas_fallback = _parse_measurement(_value_after_colon(raw))
            elif c == 1 and n.startswith('upccode') and not upc:
                upc = _value_after_colon(raw)
            elif c == 1 and n.startswith('needcart') and 'qty' in n and qty_val is None:
                qty_val = _right_value(ws, r, c, numeric=True)

    color = color or color_fallback
    meas = meas or meas_fallback
    label = f"শিট '{sheet_name}' (রো {start})"

    if not meas:
        warnings.append(f"⚠️ {label}: Carton measurement পাওয়া যায়নি — এই ব্লক স্কিপ করা হয়েছে।")
        return None, warnings
    if qty_val is None or float(qty_val) <= 0:
        warnings.append(f"⚠️ {label}: 'Need Cartoon Qty' পাওয়া যায়নি — এই ব্লক স্কিপ করা হয়েছে।")
        return None, warnings

    l, w, h, unit = meas
    if unit is None:
        unit = 'Inch'
        warnings.append(f"⚠️ {label}: measurement-এর একক (inch/cm) লেখা নেই, Inch ধরা হয়েছে — চেক করুন।")
    for name, val in (('Style', style), ('PO', po), ('Color', color), ('UPC Code', upc)):
        if not val:
            warnings.append(f"⚠️ {label}: {name} পাওয়া যায়নি।")

    return {
        'item_name': item_name,
        'ewo_no': 'N/A',
        'style_no': style or 'N/A',
        'po_no': po or 'N/A',
        'length': l, 'width': w, 'height': h,
        'ply': ply_value,
        'qty': round(float(qty_val)),
        'pack_type': upc or 'N/A',
        'reference': color or 'N/A',
        'color': '', 'size': '',
        'delivery_date': '',
        'measurement_unit': unit,
        'delivery_place_pdf': '', 'delivery_address_pdf': '',
        '_sheet': sheet_name,
        '_source_file': filename,
    }, warnings


def _read_one_file(file_stream, filename, item_name_override, manual_ply):
    wb = load_workbook(file_stream, data_only=True)
    items, warnings = [], []
    item_name = item_name_override or 'Master Carton'
    ply_value = manual_ply.strip() if manual_ply else 'N/A'

    for sn in wb.sheetnames:
        ws = wb[sn]
        if ws.sheet_state != 'visible':
            continue  # হাইড শিট — ডাটা নেওয়া হবে না
        max_rows = min(ws.max_row, 5000)
        for start, end in _find_blocks(ws, max_rows):
            item, warns = _read_block(ws, start, end, sn, filename, item_name, ply_value)
            warnings.extend(warns)
            if item:
                items.append(item)
    return items, warnings


def combine_goodearth_joefresh_files(files, item_name_override='', manual_ply=''):
    """BATCH_REGISTRY-এর কল-সিগনেচার। files: [(BytesIO, filename), ...]।
    রিটার্ন করে (line_items, warnings)। প্রতিটা লাইন-আইটেমে '_source_file'
    থাকে, তাই 'জিপ' চেকবক্স টিক দিলে ফাইল-ভিত্তিক আলাদা Excel হয়, না দিলে
    সব একসাথে।"""
    all_items, all_warnings = [], []
    for file_stream, filename in files:
        try:
            file_stream.seek(0)
            items, warns = _read_one_file(file_stream, filename, item_name_override, manual_ply)
        except Exception as e:
            all_warnings.append(f"⚠️ '{filename}' প্রসেস করতে সমস্যা হয়েছে: {e}")
            continue
        if not items:
            all_warnings.append(f"⚠️ '{filename}': কোনো ব্লক/লাইন-আইটেম পাওয়া যায়নি (পরিচিত ফরম্যাট না?)।")
        all_items.extend(items)
        all_warnings.extend(f"{w} [ফাইল: {filename}]" if filename not in w else w for w in warns)
    return all_items, all_warnings
