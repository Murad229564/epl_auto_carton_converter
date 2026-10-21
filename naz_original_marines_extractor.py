import re
from datetime import datetime
from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# NAZ Bangladesh Ltd. — Buyer: Original Marines
# 'Carton Purchase Order' (শিট: 'OM-Carton') Excel ফরম্যাট।
#   - প্রতিটা ফাইল = একটা Style। একাধিক ফাইল/শিট আপলোড করলে প্রতিটা visible
#     শিট থেকেই ডাটা নেওয়া হয়, hidden শিট বাদ। Hidden রো/কলামও বাদ।
#   - উপরের হেডার-এরিয়া থেকে (লেবেল খুঁজে, পাশের ভ্যালু):
#       'Log. Num'  -> Gmt. PO   (যেমন NAZBL-TB-26-04440)
#       'Style'     -> Gmt. Style No (সব Master Carton এবং সব Top Bottom লাইনে)
#       'Quality'   -> '5 Ply Virgin ...' থেকে Ply (পাওয়া গেলে)
#       'Req. Date' -> Delivery Date (শুধু USE_REQ_DATE_AS_DELIVERY = True হলে)
#   - মেইন টেবিল (হেডার রো: Size / Color / ... / Booking Qty (Pc)):
#       Size             -> Pack Type
#       Color            -> Reference/SKU Number (ফাঁকা থাকলে উপরের রো-র মান)
#       L / W / H (Cm)   -> Length / Width / Height
#       Booking Qty (Pc) -> Order Qty
#   - 'Top Bottom Measurement & Qty' টেবিল (হেডার: L (Cm) / W (Cm) / Qty (Pc)):
#       প্রতিটা ব্রেকডাউন রো আলাদা Top Bottom লাইন — Ply 3, Height নেই,
#       Style = ফাইলের Style, PO = Log. Num, Reference/Pack Type = 'N/A'।
#   - Item Name: UI সিলেকশন (ডিফল্ট Master Carton)।
#   - Ply (Master Carton): Quality লাইনে 'N Ply' থাকলে সেটা; না থাকলে UI-এর
#     manual_ply; তাও না থাকলে 'N/A'।
# ---------------------------------------------------------------------------

NA = 'N/A'

# Top Bottom লাইনেও PO (Log. Num) বসবে কিনা। False করলে Top Bottom-এ PO = 'N/A'।
TB_INCLUDE_PO = True

# 'Req. Date'-কে Delivery Date হিসেবে বসাতে চাইলে True করুন (ডিফল্ট: ফাঁকা,
# যাতে সিস্টেমের নিজস্ব ডেলিভারি-ডেট লজিক চলে)।
USE_REQ_DATE_AS_DELIVERY = False


def _norm(s):
    return re.sub(r'[^a-z0-9]', '', str(s or '').lower())


def _clean(v):
    if v is None:
        return ''
    return re.sub(r'\s+', ' ', str(v)).strip()


def _is_num(v):
    if v is None or isinstance(v, bool):
        return False
    try:
        float(v)
        return True
    except (TypeError, ValueError):
        return False


def _fmt_num(v):
    if v is None or v == '':
        return ''
    try:
        f = float(v)
    except (TypeError, ValueError):
        return ''
    return str(int(f)) if f == int(f) else str(round(f, 3))


def _hidden_col_indexes(ws):
    hidden = set()
    for dim in ws.column_dimensions.values():
        if dim.hidden:
            lo = dim.min or 0
            hi = dim.max or lo
            hidden.update(range(lo, hi + 1))
    return hidden


def _hidden_row_indexes(ws):
    return {r for r, dim in ws.row_dimensions.items() if dim.hidden}


def _format_date(v):
    if isinstance(v, datetime):
        return v.strftime('%d-%b-%Y')
    return _clean(v)


def _label_value(ws, label_norm, hidden_cols, max_row=15):
    """লেবেল সেল (normalized == label_norm) খুঁজে একই রো-তে ডানদিকের প্রথম
    নন-এম্পটি সেলের মান রিটার্ন করে।"""
    for r in range(1, max_row + 1):
        for c in range(1, ws.max_column + 1):
            if _norm(ws.cell(row=r, column=c).value) == label_norm:
                for c2 in range(c + 1, ws.max_column + 1):
                    if c2 in hidden_cols:
                        continue
                    v = ws.cell(row=r, column=c2).value
                    if v is not None and _clean(v) != '':
                        return v
    return None


def _find_main_header(ws, hidden_cols, max_scan=30):
    """Size + Color + Booking Qty (Pc) আছে এমন রো। হেডিং -> কলাম ম্যাপ রিটার্ন।"""
    for r in range(1, min(ws.max_row, max_scan) + 1):
        heads = {}
        for c in range(1, ws.max_column + 1):
            if c in hidden_cols:
                continue
            v = ws.cell(row=r, column=c).value
            if v is not None:
                heads.setdefault(_norm(v), c)
        if 'size' in heads and 'color' in heads and 'bookingqtypc' in heads:
            return r, heads
    return None, None


def _find_tb_header(ws, hidden_cols, start_row):
    """Top Bottom টেবিলের হেডার (L (Cm) / W (Cm) / Qty (Pc), 'Booking' ছাড়া)।"""
    for r in range(start_row, ws.max_row + 1):
        heads = {}
        for c in range(1, ws.max_column + 1):
            if c in hidden_cols:
                continue
            v = ws.cell(row=r, column=c).value
            if v is not None:
                heads.setdefault(_norm(v), c)
        if 'lcm' in heads and 'wcm' in heads and 'qtypc' in heads and 'bookingqtypc' not in heads:
            return r, heads
    return None, None


def read_naz_style_excel(file_stream, filename='', item_name_override='Master Carton', manual_ply=''):
    """মূল entry point। এই ফরম্যাট না হলে (কোনো শিটেই হেডার না পেলে) খালি
    লিস্ট [] রিটার্ন করে।"""
    wb = load_workbook(file_stream, data_only=True)
    items = []
    default_item_name = item_name_override or 'Master Carton'

    for sn in wb.sheetnames:
        ws = wb[sn]
        if ws.sheet_state != 'visible':
            continue

        hidden_cols = _hidden_col_indexes(ws)
        hidden_rows = _hidden_row_indexes(ws)

        header_row, heads = _find_main_header(ws, hidden_cols)
        if header_row is None:
            continue

        # --- উপরের হেডার-এরিয়া ---
        po_no = _clean(_label_value(ws, 'lognum', hidden_cols)) or NA
        style_no = _clean(_label_value(ws, 'style', hidden_cols)) or NA

        quality = _clean(_label_value(ws, 'quality', hidden_cols))
        m = re.search(r'(\d+)\s*ply', quality, re.I)
        if m:
            ply_value = m.group(1)
        else:
            ply_value = manual_ply.strip() if manual_ply else NA

        delivery_date = ''
        if USE_REQ_DATE_AS_DELIVERY:
            delivery_date = _format_date(_label_value(ws, 'reqdate', hidden_cols))

        col_size = heads['size']
        col_color = heads['color']
        col_l, col_w = heads.get('lcm'), heads.get('wcm')
        col_h = heads.get('hcm')
        col_qty = heads['bookingqtypc']

        def val(r, c):
            return ws.cell(row=r, column=c).value if c else None

        tb_header_row, tb_heads = _find_tb_header(ws, hidden_cols, header_row + 1)
        main_end = (tb_header_row - 1) if tb_header_row else ws.max_row

        # --- Master Carton ---
        last_color = ''
        for r in range(header_row + 1, main_end + 1):
            if r in hidden_rows:
                continue
            size_txt = _clean(val(r, col_size))
            if _norm(size_txt).startswith('total'):
                break  # মেইন টেবিল শেষ
            qty_raw = val(r, col_qty)
            if not size_txt or not _is_num(qty_raw) or float(qty_raw) <= 0:
                continue
            if not _is_num(val(r, col_l)) or not _is_num(val(r, col_w)):
                continue

            color = _clean(val(r, col_color))
            if color:
                last_color = color
            else:
                color = last_color

            items.append({
                'item_name': default_item_name,
                'ewo_no': NA,
                'style_no': style_no,
                'po_no': po_no,
                'length': _fmt_num(val(r, col_l)),
                'width': _fmt_num(val(r, col_w)),
                'height': _fmt_num(val(r, col_h)),
                'ply': ply_value,
                'qty': round(float(qty_raw)),
                'pack_type': size_txt,
                'reference': color,
                'color': '',
                'size': '',
                'delivery_date': delivery_date,
                'measurement_unit': 'Cm',
                'delivery_place_pdf': '',
                'delivery_address_pdf': '',
                '_sheet': sn,
                '_source_file': filename,
            })

        # --- Top Bottom (প্রতিটা ব্রেকডাউন আলাদা লাইন) ---
        if tb_header_row:
            tl, tw, tq = tb_heads['lcm'], tb_heads['wcm'], tb_heads['qtypc']
            for r in range(tb_header_row + 1, ws.max_row + 1):
                if r in hidden_rows:
                    continue
                # 'Total' লেবেল (L কলামে, বা মার্জড সেলের বামদিকে) পেলে টেবিল শেষ
                if any(_norm(ws.cell(row=r, column=c).value).startswith('total')
                       for c in range(1, tq + 1)):
                    break
                qty_raw = val(r, tq)
                if not _is_num(val(r, tl)) or not _is_num(val(r, tw)):
                    continue
                if not _is_num(qty_raw) or float(qty_raw) <= 0:
                    continue
                items.append({
                    'item_name': 'Top Bottom',
                    'ewo_no': NA,
                    'style_no': style_no,
                    'po_no': po_no if TB_INCLUDE_PO else NA,
                    'length': _fmt_num(val(r, tl)),
                    'width': _fmt_num(val(r, tw)),
                    'height': '',
                    'ply': '3',
                    'qty': round(float(qty_raw)),
                    'pack_type': NA,
                    'reference': NA,
                    'color': '',
                    'size': '',
                    'delivery_date': delivery_date,
                    'measurement_unit': 'Cm',
                    'delivery_place_pdf': '',
                    'delivery_address_pdf': '',
                    '_sheet': sn,
                    '_source_file': filename,
                })

    return items
