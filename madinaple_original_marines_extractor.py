import re
from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# Madinaple Fashions Craft Limited — Buyer: Original Marines
# 'QUALITY- 5 PLY BULK CARTON ORDER' Excel ফরম্যাট।
#   - একটা ফাইলে একাধিক শিট থাকে, প্রতিটা শিট একটা করে প্যাক-টাইপ (যেমন
#     '(ASSORT PACK)', '(SOLID PACK-5-6)' ইত্যাদি) — প্রতিটা শিট থেকেই
#     ডাটা নেওয়া হয়।
#   - প্রতিটা শিটে তথ্য লেবেল-ভ্যালু আকারে ছড়ানো থাকে (ফিক্সড কলাম-টেবিল
#     না), তাই লেবেল-টেক্সট খুঁজে তার পাশের প্রথম নন-এম্পটি সেলকে ভ্যালু
#     ধরা হয়:
#       'ARTICOLO : DHPF3036F'  -> Style (কোলনের পরের অংশ)
#       'M. BOX SIZE : 54X48X15CM' -> Carton Measurement (L x W x H)
#       'Carton Qty : 420'      -> Master Carton-এর Qty
#       'Top Bottom Qty : 840'  -> Top Bottom-এর Qty
#   - শিটের নাম (যেমন '(ASSORT PACK)') -> Gmt. PO কলামে বসে (ইউজার-কনফার্মড
#     — এই ফরম্যাটে আলাদা কোনো PO-নম্বর কলাম নেই, শিটের নামই সেই
#     আইডেন্টিফায়ার)।
#   - Top Bottom-এর নিজস্ব কোনো Measurement কলাম নেই — Carton Measurement
#     থেকে Length/Width প্রতিটা থেকে ৫ cm বাদ দিয়ে বের করা হয় (ইউজার-
#     কনফার্মড), Height লাগে না (Top Bottom সবসময় শুধু L x W)।
#   - এই ফরম্যাটে Divider নেই (শুধু Master Carton + Top Bottom)।
#   - Top Bottom-এর Ply সবসময় ফিক্সড ৩ (অন্য সব কাস্টমারের Top Bottom/
#     Divider কনভেনশনের সাথে সামঞ্জস্যপূর্ণ ধরে নেওয়া হয়েছে — নিশ্চিত না
#     হলে ইউজারকে জিজ্ঞেস করে নিশ্চিত হওয়া ভালো)। Master Carton-এর Ply
#     UI থেকে সিলেক্ট করা।
# ---------------------------------------------------------------------------


def _norm(s):
    return re.sub(r'[^a-z0-9]', '', str(s or '').lower())


def _clean(v):
    if v is None:
        return ''
    return re.sub(r'\s+', ' ', str(v)).strip()


def _is_num(v):
    if v is None:
        return False
    try:
        float(v)
        return True
    except (TypeError, ValueError):
        return False


def _fmt_num(v):
    if v is None or v == '':
        return ''
    try:
        f = float(v)
    except (TypeError, ValueError):
        return ''
    return str(int(f)) if f == int(f) else str(round(f, 3))


_MEAS_RE = re.compile(r'(\d+\.?\d*)\s*[xX×]\s*(\d+\.?\d*)\s*[xX×]\s*(\d+\.?\d*)')


def _parse_measurement(text):
    """'54X48X15CM' -> ('54','48','15')।"""
    if not text:
        return '', '', ''
    m = _MEAS_RE.search(str(text))
    if not m:
        return '', '', ''
    return _fmt_num(m.group(1)), _fmt_num(m.group(2)), _fmt_num(m.group(3))


def _find_label_value(ws, label_substr, max_scan=40):
    """লেবেল-টেক্সটে label_substr (normalized) থাকা সেল খুঁজে, একই রো-তে
    তার ডানদিকের প্রথম নন-এম্পটি সেলকে ভ্যালু হিসেবে রিটার্ন করে। লেবেল
    সেলেই কোলনের পরে ভ্যালু থাকলে (যেমন 'ARTICOLO : DHPF3036F' এক সেলেই)
    সেটাও ধরা হয়।"""
    for r in range(1, max_scan + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            norm_v = _norm(v)
            if label_substr not in norm_v:
                continue
            # একই সেলে কোলনের পরে ভ্যালু থাকলে (যেমন 'ARTICOLO : DHPF3036F')
            if ':' in str(v):
                after = str(v).split(':', 1)[1].strip()
                if after:
                    return after, r
            # না থাকলে একই রো-তে ডানদিকের প্রথম নন-এম্পটি, নন-পাংচুয়েশন
            # সেল (যেমন 'M. BOX SIZE' লেবেলের ঠিক পরেই একটা আলাদা ':' সেল
            # থাকে, যেটা আসল ভ্যালু না — সেটা স্কিপ করা হয়)
            for c2 in range(c + 1, ws.max_column + 1):
                v2 = ws.cell(row=r, column=c2).value
                v2_clean = _clean(v2) if v2 is not None else ''
                if v2_clean and v2_clean not in (':', '-', '.'):
                    return v2, r
    return None, None


def read_madinaple_style_excel(file_stream, filename='', item_name_override='Master Carton', manual_ply=''):
    """মূল entry point। এই ফরম্যাট না হলে (কোনো শিটেই দরকারি লেবেল না
    পেলে) খালি লিস্ট [] রিটার্ন করে।"""
    wb = load_workbook(file_stream, data_only=True)
    all_items = []
    default_item_name = item_name_override or 'Master Carton'
    ply_value = manual_ply.strip() if manual_ply else 'N/A'

    for sn in wb.sheetnames:
        ws = wb[sn]
        if ws.sheet_state != 'visible':
            continue  # হাইড/ভেরি-হাইড শীট — ডাটা নেওয়া হবে না

        style_val, _ = _find_label_value(ws, 'articolo')
        box_size_val, _ = _find_label_value(ws, 'mboxsize')
        carton_qty_val, _ = _find_label_value(ws, 'cartonqty')
        tb_qty_val, _ = _find_label_value(ws, 'topbottomqty')

        if style_val is None or box_size_val is None:
            continue  # এই শিট এই ফরম্যাটের না (প্রত্যাশিত লেবেল পাওয়া যায়নি)

        style_no = _clean(style_val) or 'N/A'
        po_no = _clean(sn) or 'N/A'
        c_l, c_w, c_h = _parse_measurement(box_size_val)

        if c_l and _is_num(carton_qty_val) and float(carton_qty_val) > 0:
            all_items.append({
                'item_name': default_item_name,
                'ewo_no': 'N/A',
                'style_no': style_no,
                'po_no': po_no,
                'length': c_l,
                'width': c_w,
                'height': c_h,
                'ply': ply_value,
                'qty': round(float(carton_qty_val)),
                'pack_type': '',
                'reference': '',
                'color': '',
                'size': '',
                'delivery_date': '',
                'measurement_unit': 'Cm',
                'delivery_place_pdf': '',
                'delivery_address_pdf': '',
                '_sheet': sn,
                '_source_file': filename,
            })

        if c_l and c_w and _is_num(tb_qty_val) and float(tb_qty_val) > 0:
            tb_l = _fmt_num(float(c_l) - 5)
            tb_w = _fmt_num(float(c_w) - 5)
            all_items.append({
                'item_name': 'Top Bottom',
                'ewo_no': 'N/A',
                'style_no': style_no,
                'po_no': po_no,
                'length': tb_l,
                'width': tb_w,
                'height': '',
                'ply': '3',
                'qty': round(float(tb_qty_val)),
                'pack_type': '',
                'reference': '',
                'color': '',
                'size': '',
                'delivery_date': '',
                'measurement_unit': 'Cm',
                'delivery_place_pdf': '',
                'delivery_address_pdf': '',
                '_sheet': sn,
                '_source_file': filename,
            })

    return all_items
