import re
from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# Knit Asia Ltd. — Buyer: Kohl`s
# 'Carton Measurement & Booking Sheet' এক্সেল ফরম্যাট।
#   - হেডার দুই রো জুড়ে থাকে:
#       রো ১ (গ্রুপ-লেবেল): IR / Style / PO Number / PO type / Category /
#         Qty / Cartons Measurement (গ্রুপ) / CBM / Poly measurement (গ্রুপ) /
#         Remarks
#       রো ২ (সাব-লেবেল, 'Cartons Measurement' গ্রুপের নিচে): L / W / H /
#         Per ctn pcs/Set / Req. Ctn
#   - ইউজার-কনফার্মড ম্যাপিং:
#       IR          -> Remarks (Template-এর Remarks কলামে)
#       Style       -> Gmt. Style No
#       PO Number   -> Gmt. PO
#       PO type     -> Reference/SKU Number
#       Category    -> Pack Type
#       L/W/H       -> Length/Width/Height
#       Req. Ctn    -> Order Qty (এই ফরম্যাটে 'Qty' কলামের টোটাল-পিস ভ্যালু
#                      ব্যবহার হয় না, শুধু Req. Ctn — required carton সংখ্যা)
#       Ply         -> সবসময় ফিক্সড ৫ (ইউজার-কনফার্মড)
#   - Item Name সাধারণত UI থেকে সিলেক্ট করা ভ্যালু (item_name_override)
#     বসে, কিন্তু সোর্স Excel-এর নিজস্ব 'Remarks' কলামে (IR কলাম না — এটা
#     আলাদা, ডানদিকে থাকা আসল 'Remarks' হেডার-ওয়ালা কলাম) 'ELASTIC' শব্দ
#     থাকলে সেই রো-তে Item Name 'Elastic Hanger Carton' হয়ে যায়
#     (ওভাররাইড, UI সিলেকশন যাই হোক না কেন)।
#   - 'Sub Total'/'Total' রো এবং L/W/H ফাঁকা এমন রো (ব্লকের মাঝে খালি
#     সেপারেটর রো, #DIV/0! ফর্মুলা-এরর সহ) — এসব বাদ যায়, শুধু L/W/H
#     আছে এমন রো-ই ডাটা হিসেবে নেওয়া হয়।
#   - Excel-এ hidden রো/কলাম থাকলে সেগুলো একদম ব্যবহার হয় না — hidden
#     কলামে কোনো ফিল্ড ম্যাপ হলেও সেটা ইগনোর করা হয়, hidden রো ডাটা
#     হিসেবেই নেওয়া হয় না।
#   - একই ফাইলে একাধিক শিট থাকতে পারে — প্রতিটা শিট থেকেই ডাটা নেওয়া হয়।
# ---------------------------------------------------------------------------


def _norm(s):
    return re.sub(r'[^a-z0-9]', '', str(s or '').lower())


def _clean(v):
    if v is None:
        return ''
    return re.sub(r'\s+', ' ', str(v)).strip()


def _is_num(v):
    if v is None:
        return False
    try:
        float(v)
        return True
    except (TypeError, ValueError):
        return False


def _fmt_num(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return ''
    return str(int(f)) if f == int(f) else str(round(f, 3))


def _get_hidden_rows_cols(ws):
    """Excel-এ hidden করা রো এবং কলামের (1-indexed) সেট রিটার্ন করে।"""
    hidden_rows = {r for r, dim in ws.row_dimensions.items() if dim.hidden}
    from openpyxl.utils import column_index_from_string
    hidden_cols = set()
    for letter, dim in ws.column_dimensions.items():
        if not dim.hidden:
            continue
        try:
            hidden_cols.add(column_index_from_string(letter))
        except Exception:
            pass
    return hidden_rows, hidden_cols


def _find_header_row(ws, max_scan=20):
    """'IR'-এর ঠিক পরের কলামে ' Style'/'Style' (leading space সহ্য করে)
    এই কম্বিনেশন খুঁজে হেডার রো (গ্রুপ-লেবেল রো) বের করে।"""
    for r in range(1, max_scan + 1):
        for c in range(1, ws.max_column + 1):
            if _norm(ws.cell(row=r, column=c).value) == 'ir':
                nxt = _norm(ws.cell(row=r, column=c + 1).value)
                if nxt.startswith('style'):
                    return r
    return None


def _build_col_map(ws, header_row, hidden_cols):
    """হেডার রো (গ্রুপ-লেবেল) আর তার ঠিক পরের রো (সাব-লেবেল: L/W/H/Req. Ctn)
    দুটোই স্ক্যান করে কলাম-ম্যাপ বানায়। hidden কলামে কোনো লেবেল মিললেও
    সেটা ব্যবহার হয় না (ইউজার-কনফার্মড নিয়ম)।"""
    col_map = {}
    n_cols = ws.max_column

    def maybe_set(key, c):
        if c not in hidden_cols and key not in col_map:
            col_map[key] = c

    for c in range(1, n_cols + 1):
        label = _norm(ws.cell(row=header_row, column=c).value)
        if not label:
            continue
        if label == 'ir':
            maybe_set('ir_remarks', c)
        elif label.startswith('style'):
            maybe_set('style_no', c)
        elif label == 'ponumber':
            maybe_set('po_no', c)
        elif label == 'potype':
            maybe_set('reference', c)
        elif label == 'category':
            maybe_set('pack_type', c)
        elif label == 'remarks':
            maybe_set('src_remarks', c)  # সোর্সের নিজস্ব Remarks — শুধু ELASTIC ডিটেক্ট করতে

    sub_row = header_row + 1
    for c in range(1, n_cols + 1):
        label = _norm(ws.cell(row=sub_row, column=c).value)
        if not label:
            continue
        if label == 'l':
            maybe_set('length', c)
        elif label == 'w':
            maybe_set('width', c)
        elif label == 'h':
            maybe_set('height', c)
        elif 'reqctn' in label:
            maybe_set('qty', c)

    return col_map


def read_knitasia_style_excel(file_stream, filename='', item_name_override='Master Carton', manual_ply=''):
    """মূল entry point। এই ফরম্যাট না হলে (হেডার না মিললে) খালি লিস্ট []
    রিটার্ন করে।"""
    wb = load_workbook(file_stream, data_only=True)
    all_items = []
    default_item_name = item_name_override or 'Master Carton'

    for sn in wb.sheetnames:
        ws = wb[sn]
        if ws.sheet_state != 'visible':
            continue  # হাইড/ভেরি-হাইড শীট — এই শীটের কোনো ডাটাই নেওয়া হবে না
        header_row = _find_header_row(ws)
        if header_row is None:
            continue

        hidden_rows, hidden_cols = _get_hidden_rows_cols(ws)
        col_map = _build_col_map(ws, header_row, hidden_cols)

        required = ('style_no', 'po_no', 'length', 'width', 'height', 'qty')
        if not all(k in col_map for k in required):
            continue  # প্রত্যাশিত কলাম পাওয়া যায়নি (বা hidden) — এই ফরম্যাট না

        ir_col = col_map.get('ir_remarks')
        style_col = col_map['style_no']
        po_col = col_map['po_no']
        ref_col = col_map.get('reference')
        pack_col = col_map.get('pack_type')
        len_col = col_map['length']
        wid_col = col_map['width']
        hgt_col = col_map['height']
        qty_col = col_map['qty']
        src_remarks_col = col_map.get('src_remarks')

        r = header_row + 2  # গ্রুপ-লেবেল রো + সাব-লেবেল রো, তারপরই ডাটা শুরু
        max_row = ws.max_row
        while r <= max_row:
            if r in hidden_rows:
                r += 1
                continue

            first_cell = _clean(ws.cell(row=r, column=1).value)
            if _norm(first_cell) in ('subtotal', 'grandtotal', 'total'):
                r += 1
                continue

            length_v = ws.cell(row=r, column=len_col).value
            width_v = ws.cell(row=r, column=wid_col).value
            height_v = ws.cell(row=r, column=hgt_col).value
            if not (_is_num(length_v) and _is_num(width_v) and _is_num(height_v)):
                r += 1
                continue  # L/W/H না থাকা মানে এটা সাব-টোটাল/খালি সেপারেটর রো

            style_val = _clean(ws.cell(row=r, column=style_col).value)
            po_val = ws.cell(row=r, column=po_col).value
            qty_val = ws.cell(row=r, column=qty_col).value
            if not style_val or not _is_num(qty_val):
                r += 1
                continue

            item_name = default_item_name
            if src_remarks_col is not None:
                remark_text = str(ws.cell(row=r, column=src_remarks_col).value or '').lower()
                if 'elastic' in remark_text:
                    item_name = 'Elastic Hanger Carton'

            po_str = str(po_val).strip() if isinstance(po_val, (int, float)) and float(po_val).is_integer() \
                else _clean(po_val)
            if isinstance(po_val, (int, float)):
                po_str = str(int(po_val)) if float(po_val).is_integer() else str(po_val)

            all_items.append({
                'item_name': item_name,
                'ewo_no': 'N/A',
                'style_no': style_val,
                'po_no': po_str,
                'length': _fmt_num(length_v),
                'width': _fmt_num(width_v),
                'height': _fmt_num(height_v),
                'ply': '5',
                'qty': qty_val,
                'pack_type': _clean(ws.cell(row=r, column=pack_col).value) if pack_col else '',
                'reference': _clean(ws.cell(row=r, column=ref_col).value) if ref_col else '',
                'remarks': _clean(ws.cell(row=r, column=ir_col).value) if ir_col else '',
                'color': '',
                'size': '',
                'delivery_date': '',
                'measurement_unit': 'Inch',
                'delivery_place_pdf': '',
                'delivery_address_pdf': '',
                '_sheet': sn,
                '_source_file': filename,
            })
            r += 1

    return all_items
