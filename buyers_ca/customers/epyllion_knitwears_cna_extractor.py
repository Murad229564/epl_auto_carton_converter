"""
Epyllion Knitwears Limited (IN-HOUSE) — Buyer: C&A
'Purchase Order (PO)' মাল্টি-শিট Excel ফরম্যাট।
  - Sheet1: টোটাল সামারি (Rate/Total Value সহ) — শুধু ইউজারের নিজের
    ম্যানুয়াল প্রাইস-মিলানোর জন্য, extractor এখান থেকে কিছু নেয় না।
  - Sheet2 ('Purchase Order Details'): আসল ব্রেকডাউন, এখান থেকেই ডাটা
    নেওয়া হয়:
        EWO No       -> Gmt. EWO No
        Style No     -> Style (raw, /TRI ca_rules.py-তে যোগ হয়)
        Ply          -> Ply (কলামে থাকলে সেটাই, ফাইলে সবসময় থাকে)
        Measurement  -> L/W/H (যেমন 'L- 50 X  W- 30 X  H- 10 cm', একক
                        সবসময় cm)
        Instruction  -> raw code-text (C&A কোড থাকলে সেটাই, 'Irregular'
                        হলে বা ফাঁকা হলে ca_rules.py নিজেই এক্সেপশনাল
                        ধরে নেয়)
        QTY          -> Qty
  - 'PCS Wise Total' রো-তে থেমে যায়। hidden রো থাকলে বাদ (দেখা যায়নি
    এই ফরম্যাটের Sheet2-তে, কিন্তু ভবিষ্যতের জন্য চেক রাখা হয়েছে)।
"""
import re
from openpyxl import load_workbook

from ..ca_rules import apply_ca_rules


def _norm(s):
    return re.sub(r'[^a-z0-9]', '', str(s or '').lower())


def _clean(v):
    if v is None:
        return ''
    return re.sub(r'\s+', ' ', str(v)).strip()


_MEAS_RE = re.compile(r'L[\s\-]*(\d+\.?\d*).*?W[\s\-]*(\d+\.?\d*).*?H[\s\-]*(\d+\.?\d*)', re.I | re.S)


def _parse_measurement_cm(text):
    if not text:
        return None, None, None
    m = _MEAS_RE.search(str(text))
    if not m:
        return None, None, None
    return float(m.group(1)), float(m.group(2)), float(m.group(3))


def _find_detail_sheet(wb):
    """'Purchase Order Details' লেখা টাইটেল-সহ শিটটা খুঁজে বের করে, শিটের
    নাম ('Sheet2' ইত্যাদি) নির্বিশেষে — যাতে ফাইল-ভেদে শিট-অর্ডার বদলে
    গেলেও কাজ করে।"""
    for sn in wb.sheetnames:
        ws = wb[sn]
        if ws.sheet_state != 'visible':
            continue
        v = ws.cell(row=1, column=1).value
        if v and 'purchase order details' in str(v).lower():
            return ws, sn
    return None, None


def read_epyllion_cna_excel(file_stream, filename='', lookup=None, item_name_override='', manual_ply=''):
    """মূল entry point। রিটার্ন করে (line_items, warnings)।"""
    wb = load_workbook(file_stream, data_only=True)
    ws, sn = _find_detail_sheet(wb)
    if ws is None:
        return [], []

    header_row = None
    col_map = {}
    for r in range(1, min(ws.max_row, 5) + 1):
        row_labels = {}
        for c in range(1, ws.max_column + 1):
            label = _norm(ws.cell(row=r, column=c).value)
            if label:
                row_labels[label] = c
        if 'ewono' in row_labels and 'styleno' in row_labels:
            header_row = r
            col_map = row_labels
            break
    if header_row is None:
        return [], []

    ewo_col = col_map.get('ewono')
    style_col = col_map.get('styleno')
    ply_col = col_map.get('ply')
    meas_col = col_map.get('measurement')
    instr_col = col_map.get('instruction')
    qty_col = col_map.get('qty')
    if not all([ewo_col, style_col, meas_col, qty_col]):
        return [], [f"⚠️ '{filename}': প্রত্যাশিত কলাম (EWO No/Style No/Measurement/QTY) পাওয়া যায়নি।"]

    hidden_rows = {r for r, dim in ws.row_dimensions.items() if dim.hidden}

    line_items = []
    all_warnings = []
    r = header_row + 1
    max_row = ws.max_row
    while r <= max_row:
        if r in hidden_rows:
            r += 1
            continue

        first_col_text = _norm(ws.cell(row=r, column=1).value)
        if 'pcswisetotal' in first_col_text or first_col_text == 'total':
            break

        style_val = _clean(ws.cell(row=r, column=style_col).value)
        qty_val = ws.cell(row=r, column=qty_col).value
        if not style_val or qty_val is None:
            r += 1
            continue

        l_cm, w_cm, h_cm = _parse_measurement_cm(ws.cell(row=r, column=meas_col).value)
        if l_cm is None:
            r += 1
            continue

        raw_item = {
            'raw_style': style_val,
            'ewo_no': _clean(ws.cell(row=r, column=ewo_col).value),
            'raw_code_text': _clean(ws.cell(row=r, column=instr_col).value) if instr_col else '',
            'length_cm': l_cm, 'width_cm': w_cm, 'height_cm': h_cm,
            'qty': qty_val,
            'ply': _clean(ws.cell(row=r, column=ply_col).value) if ply_col else '',
            '_sheet': sn, '_source_file': filename,
        }
        item, warns = apply_ca_rules(raw_item, lookup, item_name_override=item_name_override, manual_ply=manual_ply)
        if item:
            line_items.append(item)
        all_warnings.extend(warns)
        r += 1

    return line_items, all_warnings
