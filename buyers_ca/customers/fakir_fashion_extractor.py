"""
Fakir Fashion Limited — Buyer: C&A
'C&A Carton booking -...' Excel ফরম্যাট।
  - একটা শিটে একাধিক 'Order No :' ব্লক থাকে (প্রতিটা ব্লকের নিজের
    হেডার-রো: Item code | Measurement (cm): | Quantity (pcs) | REMARKS)।
  - Style: 'Order No :'-এর পরের টেক্সট (colon-এর পর যা আছে) — /TRI যোগ
    হয় ca_rules.py-তেই, এখানে raw টেক্সট-ই পাঠানো হয়।
  - Item code কলাম ফাঁকা বা 'Regular'-জাতীয় (C&A-প্যাটার্নের বাইরে)
    হলে ca_rules.py নিজেই এক্সেপশনাল ধরে নেয় — এখানে আলাদা চেক লাগে না।
  - 'G.Total' রো-তে থেমে যায়।
"""
import re
from openpyxl import load_workbook

from ..ca_rules import apply_ca_rules


def _norm(s):
    return re.sub(r'[^a-z0-9]', '', str(s or '').lower())


def _clean(v):
    if v is None:
        return ''
    return re.sub(r'\s+', ' ', str(v)).strip()


_MEAS_RE = re.compile(r'(\d+\.?\d*)\s*[xX×]\s*(\d+\.?\d*)\s*[xX×]\s*(\d+\.?\d*)')


def _parse_measurement_cm(text):
    if not text:
        return None, None, None
    m = _MEAS_RE.search(str(text))
    if not m:
        return None, None, None
    return float(m.group(1)), float(m.group(2)), float(m.group(3))


_ORDER_NO_RE = re.compile(r'order\s*no\s*:?\s*(.+)', re.I)


def read_fakir_fashion_excel(file_stream, filename='', lookup=None, item_name_override='', manual_ply=''):
    """মূল entry point। রিটার্ন করে (line_items, warnings)। lookup
    (ca_calc.build_lookup()-এর রেজাল্ট) caller থেকে পাস করতে হবে।"""
    wb = load_workbook(file_stream, data_only=True)
    line_items = []
    all_warnings = []

    for sn in wb.sheetnames:
        ws = wb[sn]
        if ws.sheet_state != 'visible':
            continue

        header_row = None
        for r in range(1, min(ws.max_row, 15) + 1):
            b = _norm(ws.cell(row=r, column=2).value)
            if b == 'itemcode':
                header_row = r
                break
        if header_row is None:
            continue  # এই শিট এই ফরম্যাটের না

        current_order = ''
        r = 1
        max_row = ws.max_row
        consecutive_blank = 0
        while r <= max_row:
            b_val = ws.cell(row=r, column=2).value
            b_clean = _clean(b_val)
            b_norm = _norm(b_val)

            if _norm(b_val).startswith('orderno'):
                m = _ORDER_NO_RE.search(str(b_val))
                if m:
                    current_order = m.group(1).strip()
                r += 1
                continue
            if b_norm == 'itemcode':
                r += 1
                continue  # সাব-হেডার রো
            if b_norm.startswith('gtotal') or b_norm == 'total':
                break

            if not b_clean and not _clean(ws.cell(row=r, column=3).value):
                consecutive_blank += 1
                if consecutive_blank >= 50:
                    break
                r += 1
                continue
            consecutive_blank = 0

            meas_text = ws.cell(row=r, column=3).value
            qty_val = ws.cell(row=r, column=4).value
            l_cm, w_cm, h_cm = _parse_measurement_cm(meas_text)
            if l_cm is None or qty_val is None:
                r += 1
                continue

            raw_item = {
                'raw_style': current_order,
                'raw_code_text': b_clean,
                'length_cm': l_cm, 'width_cm': w_cm, 'height_cm': h_cm,
                'qty': qty_val,
                '_sheet': sn, '_source_file': filename,
            }
            item, warns = apply_ca_rules(raw_item, lookup, item_name_override=item_name_override, manual_ply=manual_ply)
            if item:
                line_items.append(item)
            all_warnings.extend(warns)
            r += 1

    return line_items, all_warnings
