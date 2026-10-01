import re
from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# Green Life Knit Composite Ltd. — Buyer: Abode Sourcing Ltd.
# 'BULK CARTON WORK ORDER' Excel ফরম্যাট।
#   - একটা ফাইলে একাধিক শিট থাকতে পারে, প্রতিটা শিট সাধারণত একটা স্টাইলের
#     বুকিং (শিটের নামও স্টাইল কোড অনুযায়ী, যেমন 'AQ010', 'AQ011')।
#   - হেডার রো (সাধারণত ২৪ নম্বরে, কিন্তু fixed ধরা হয় না): STYLE COLOUR |
#     P.ORDER No | SIZE | QUANTITY | Measurement | CARTON QTY | SQM Price |
#     Unit Price | Total Value।
#   - ইউজার-কনফার্মড ম্যাপিং:
#       STYLE COLOUR -> Gmt. Style No
#       P.ORDER No   -> Gmt. PO
#       SIZE         -> Reference/SKU Number
#       Measurement  -> Length/Width/Height
#       CARTON QTY   -> Order Qty (QUANTITY কলাম — যেটাতে '36 PCS'-এর
#                       মতো থাকে — ব্যবহার হয় না)
#   - প্রতিটা স্টাইলের সাইজ-ব্রেকডাউনের পর দুইটা আলাদা সামারি-লাইন থাকে:
#       'Top/Bottom' রো -> Item Name 'Top Bottom', Ply 3
#       'Divider' রো     -> Item Name 'Divider', Ply 3
#     দুটোরই নিজস্ব Measurement/Qty থাকে (Style/PO/Reference প্রযোজ্য না,
#     তাই 'N/A')। এই দুই লাইনের Ply সবসময় ফিক্সড ৩ — সাধারণ Master Carton-এর
#     Ply (UI থেকে সিলেক্ট করা, ৩ বা ৫) থেকে আলাদা, ইউজার-কনফার্মড।
#   - এর মাঝের একটা সাব-টোটাল রো (শুধু Qty-এর যোগফল, Style/Measurement
#     ফাঁকা) আর সবশেষে 'Grand Total' রো — এই দুটোই বাদ (ডাটা না)।
# ---------------------------------------------------------------------------


def _norm(s):
    return re.sub(r'[^a-z0-9]', '', str(s or '').lower())


def _clean(v):
    if v is None:
        return ''
    return re.sub(r'\s+', ' ', str(v)).strip()


def _is_num(v):
    if v is None:
        return False
    try:
        float(v)
        return True
    except (TypeError, ValueError):
        return False


def _fmt_num(v):
    if v is None or v == '':
        return ''
    try:
        f = float(v)
    except (TypeError, ValueError):
        return ''
    return str(int(f)) if f == int(f) else str(round(f, 3))


_MEAS_RE = re.compile(r'(\d+\.?\d*)\s*[xX×]\s*(\d+\.?\d*)(?:\s*[xX×]\s*(\d+\.?\d*))?')


def _parse_measurement(text):
    """'56 x 38 x 20 CM' -> ('56','38','20')। '48 X 32 CM' (Height ছাড়া,
    Top/Bottom-Divider-এ) -> ('48','32','')।"""
    if not text:
        return '', '', ''
    m = _MEAS_RE.search(str(text))
    if not m:
        return '', '', ''
    l = _fmt_num(m.group(1))
    w = _fmt_num(m.group(2))
    h = _fmt_num(m.group(3)) if m.group(3) else ''
    return l, w, h


def _find_header_row(ws, max_scan=40):
    """কলাম A-তে 'STYLE COLOUR' আর কলাম B-তে 'P.ORDER No' — এই দুটো
    একসাথে থাকা রো-কেই হেডার ধরা হয়।"""
    for r in range(1, max_scan + 1):
        a = _norm(ws.cell(row=r, column=1).value)
        b = _norm(ws.cell(row=r, column=2).value)
        if a == 'stylecolour' and b.startswith('porderno'):
            return r
    return None


def _build_col_map(ws, header_row):
    col_map = {}
    for c in range(1, ws.max_column + 1):
        label = _norm(ws.cell(row=header_row, column=c).value)
        if not label:
            continue
        if label == 'stylecolour':
            col_map['style_no'] = c
        elif label.startswith('porderno'):
            col_map['po_no'] = c
        elif label == 'size':
            col_map['reference'] = c
        elif label == 'measurement':
            col_map['measurement'] = c
        elif 'cartonqty' in label:
            col_map['qty'] = c
    return col_map


def read_greenlife_abode_style_excel(file_stream, filename='', item_name_override='Master Carton', manual_ply=''):
    """মূল entry point। এই ফরম্যাট না হলে (কোনো শিটেই হেডার না মিললে)
    খালি লিস্ট [] রিটার্ন করে।"""
    wb = load_workbook(file_stream, data_only=True)
    all_items = []
    default_item_name = item_name_override or 'Master Carton'
    ply_value = manual_ply.strip() if manual_ply else 'N/A'

    for sn in wb.sheetnames:
        ws = wb[sn]
        if ws.sheet_state != 'visible':
            continue  # হাইড/ভেরি-হাইড শীট — ডাটা নেওয়া হবে না

        header_row = _find_header_row(ws)
        if header_row is None:
            continue

        col_map = _build_col_map(ws, header_row)
        required = ('style_no', 'po_no', 'reference', 'measurement', 'qty')
        if not all(k in col_map for k in required):
            continue  # প্রত্যাশিত কলাম পাওয়া যায়নি — এই ফরম্যাট না

        hidden_rows = {r for r, dim in ws.row_dimensions.items() if dim.hidden}
        style_col = col_map['style_no']
        po_col = col_map['po_no']
        ref_col = col_map['reference']
        meas_col = col_map['measurement']
        qty_col = col_map['qty']

        r = header_row + 1
        max_row = ws.max_row
        consecutive_blank = 0
        while r <= max_row:
            if r in hidden_rows:
                r += 1
                continue

            first_text = _clean(ws.cell(row=r, column=1).value)
            first_norm = _norm(first_text)

            if 'grandtotal' in first_norm:
                break  # টেবিলের একদম শেষ

            if not first_text:
                # ফাঁকা রো অথবা সাব-টোটাল রো (শুধু Qty-এর যোগফল, কোনো
                # Style/Item-ইনফো নেই) — ডাটা না, স্কিপ
                consecutive_blank += 1
                if consecutive_blank >= 100:
                    break
                r += 1
                continue
            consecutive_blank = 0

            qty_val = ws.cell(row=r, column=qty_col).value
            if not _is_num(qty_val) or float(qty_val) <= 0:
                r += 1
                continue

            length, width, height = _parse_measurement(ws.cell(row=r, column=meas_col).value)
            if not length:
                r += 1
                continue

            if 'top' in first_norm and 'bottom' in first_norm:
                item_name, ply, style_no, po_no, reference = 'Top Bottom', '3', 'N/A', 'N/A', 'N/A'
            elif first_norm == 'divider':
                item_name, ply, style_no, po_no, reference = 'Divider', '3', 'N/A', 'N/A', 'N/A'
            else:
                item_name = default_item_name
                ply = ply_value
                style_no = _clean(ws.cell(row=r, column=style_col).value) or 'N/A'
                po_no = _clean(ws.cell(row=r, column=po_col).value) or 'N/A'
                reference = _clean(ws.cell(row=r, column=ref_col).value) or 'N/A'

            all_items.append({
                'item_name': item_name,
                'ewo_no': 'N/A',
                'style_no': style_no,
                'po_no': po_no,
                'length': length,
                'width': width,
                'height': height,
                'ply': ply,
                'qty': round(float(qty_val)),
                'pack_type': '',
                'reference': reference,
                'color': '',
                'size': '',
                'delivery_date': '',
                'measurement_unit': 'Cm',
                'delivery_place_pdf': '',
                'delivery_address_pdf': '',
                '_sheet': sn,
                '_source_file': filename,
            })
            r += 1

    return all_items
