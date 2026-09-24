import re
from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# Fame Apparels Limited — Buyer: Defacto
# 'Summary Table-English Format' Carton বুকিং এক্সেল ফরম্যাট।
#   - টেবিলের উপরে একটা ইনফো-ব্লকে 'STYLE NAME' লেবেলের পাশে একটা
#     সিঙ্গেল স্টাইল নম্বর থাকে (যেমন 'I5771AX') — পুরো শিটের সব লাইনের
#     Gmt. Style No-তে এটাই বসে (প্রতি রো-তে আলাদা Style কলাম নেই)।
#   - মূল ডাটা-টেবিলের হেডার: Order Number / Prepack Code / [সাইজ কলামগুলো,
#     যেমন XS/S/M/L/XL — সংখ্যা/নামে ভিন্ন হতে পারে] / TTL QTY / CARTON QTY /
#     CARTON MESURMENT (বানান-ভুল সহ্য করে ম্যাচ করা হয়)।
#   - CARTON MESURMENT কলামটা merged cell (শুধু প্রথম ডাটা-রো-তে ভ্যালু
#     দেখায়, openpyxl বাকি merged রো-গুলোতে None রিটার্ন করে) — তাই
#     forward-fill করে সব রো-তে একই measurement বসানো হয়।
#   - ইউজার-কনফার্মড ম্যাপিং:
#       STYLE NAME (গ্লোবাল)  -> Gmt. Style No (সব রো-তে একই)
#       Order Number          -> Gmt. PO
#       Prepack Code          -> Reference/SKU Number
#       CARTON QTY            -> Order Qty (মূল কোয়ান্টিটি — TTL QTY/সাইজ
#                                 কলামগুলো ব্যবহার হয় না)
#       CARTON MESURMENT      -> Length/Width/Height
#       Item Name/Ply         -> UI থেকে সিলেক্ট করা ভ্যালু (item_name_override/
#                                 manual_ply), ফাইলে এসবের কোনো কলাম নেই
#   - সবার শেষের 'Total' রো (Order Number ফাঁকা, শুধু সাইজ-কলাম/TTL QTY-এর
#     যোগফল) — Order Number ফাঁকা থাকায় এমনিতেই বাদ পড়ে যায়।
# ---------------------------------------------------------------------------


def _norm(s):
    return re.sub(r'[^a-z0-9]', '', str(s or '').lower())


def _clean(v):
    if v is None:
        return ''
    return re.sub(r'\s+', ' ', str(v)).strip()


def _is_num(v):
    if v is None:
        return False
    try:
        float(v)
        return True
    except (TypeError, ValueError):
        return False


def _fmt_num(v):
    if v is None or v == '':
        return ''
    try:
        f = float(v)
    except (TypeError, ValueError):
        return ''
    return str(int(f)) if f == int(f) else str(round(f, 3))


def _po_str(v):
    if isinstance(v, (int, float)):
        return str(int(v)) if float(v).is_integer() else str(v)
    return _clean(v)


_MEAS_RE = re.compile(r'(\d+\.?\d*)\s*[xX×]\s*(\d+\.?\d*)\s*[xX×]\s*(\d+\.?\d*)')


def _parse_measurement(text):
    """'60X40X35CM' -> ('60','40','35')।"""
    if not text:
        return '', '', ''
    m = _MEAS_RE.search(str(text))
    if not m:
        return '', '', ''
    return _fmt_num(m.group(1)), _fmt_num(m.group(2)), _fmt_num(m.group(3))


def _extract_style_no(ws, max_scan=10):
    """উপরের ইনফো-ব্লকে 'STYLE NAME' লেবেলের পাশে (একই রো-তে, ডানদিকের
    প্রথম নন-এম্পটি সেল) সিঙ্গেল স্টাইল নম্বরটা খুঁজে বের করে।"""
    for r in range(1, max_scan + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None and _norm(v) == 'stylename':
                for c2 in range(c + 1, ws.max_column + 1):
                    v2 = ws.cell(row=r, column=c2).value
                    if v2 is not None and _clean(v2):
                        return _clean(v2)
    return ''


def _find_header_row(ws, max_scan=20):
    """কলাম A-তে 'Order Number' আর কলাম B-তে 'Prepack Code' — এই দুটো
    একসাথে থাকা রো-কেই হেডার ধরা হয়।"""
    for r in range(1, max_scan + 1):
        a = _norm(ws.cell(row=r, column=1).value)
        b = _norm(ws.cell(row=r, column=2).value)
        if a == 'ordernumber' and b.startswith('prepack'):
            return r
    return None


def _build_col_map(ws, header_row):
    col_map = {}
    for c in range(1, ws.max_column + 1):
        label = _norm(ws.cell(row=header_row, column=c).value)
        if not label:
            continue
        if label == 'ordernumber':
            col_map['po_no'] = c
        elif label.startswith('prepack'):
            col_map['reference'] = c
        elif label == 'cartonqty':
            col_map['qty'] = c
        elif 'carton' in label and ('measur' in label or 'mesur' in label):
            col_map['measurement'] = c
    return col_map


def read_fame_defacto_style_excel(file_stream, filename='', item_name_override='Master Carton', manual_ply=''):
    """মূল entry point। এই ফরম্যাট না হলে (হেডার না মিললে) খালি লিস্ট []
    রিটার্ন করে।"""
    wb = load_workbook(file_stream, data_only=True)
    all_items = []
    default_item_name = item_name_override or 'Master Carton'
    ply_value = manual_ply.strip() if manual_ply else 'N/A'

    for sn in wb.sheetnames:
        ws = wb[sn]
        header_row = _find_header_row(ws)
        if header_row is None:
            continue

        col_map = _build_col_map(ws, header_row)
        required = ('po_no', 'reference', 'qty', 'measurement')
        if not all(k in col_map for k in required):
            continue  # প্রত্যাশিত কলাম পাওয়া যায়নি — এই ফরম্যাট না

        style_no = _extract_style_no(ws) or 'N/A'
        po_col = col_map['po_no']
        ref_col = col_map['reference']
        qty_col = col_map['qty']
        meas_col = col_map['measurement']

        r = header_row + 1
        max_row = ws.max_row
        last_measurement = ''
        while r <= max_row:
            po_val = ws.cell(row=r, column=po_col).value
            if po_val is None or not _clean(po_val):
                r += 1
                continue  # Order Number ফাঁকা — Total রো বা খালি রো, স্কিপ

            qty_val = ws.cell(row=r, column=qty_col).value
            if not _is_num(qty_val):
                r += 1
                continue

            meas_raw = ws.cell(row=r, column=meas_col).value
            if meas_raw is not None and _clean(meas_raw):
                last_measurement = str(meas_raw)
            length, width, height = _parse_measurement(last_measurement)
            if not length:
                r += 1
                continue

            all_items.append({
                'item_name': default_item_name,
                'ewo_no': 'N/A',
                'style_no': style_no,
                'po_no': _po_str(po_val),
                'length': length,
                'width': width,
                'height': height,
                'ply': ply_value,
                'qty': round(float(qty_val)) if float(qty_val).is_integer() else qty_val,
                'pack_type': '',
                'reference': _clean(ws.cell(row=r, column=ref_col).value),
                'color': '',
                'size': '',
                'delivery_date': '',
                'measurement_unit': 'Cm',
                'delivery_place_pdf': '',
                'delivery_address_pdf': '',
                '_sheet': sn,
                '_source_file': filename,
            })
            r += 1

    return all_items
