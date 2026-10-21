import re
from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# Southern Clothings Ltd. — Buyer: Original Marines
# 'Carton Booking Sheet' Excel ফরম্যাট।
#   - একটা ফাইলে একাধিক শিট থাকতে পারে (যেমন 'Regular qty', 'LIFCO qty') —
#     প্রতিটা visible শিট থেকেই ডাটা নেওয়া হয়। Hidden শিট বাদ।
#   - হেডার রো (সাধারণত রো 7) 'Style Number' + 'Carton Booking Qty' দিয়ে
#     খুঁজে বের করা হয়, কলামের অবস্থান নয় — হেডিং-নাম দিয়ে কলাম ম্যাপ হয়,
#     তাই দ্বিতীয় শিটে কলাম আগে-পিছে/কম-বেশি থাকলেও সমস্যা নেই।
#   - Excel-এ HIDDEN করা কলামের ডাটা সম্পূর্ণ উপেক্ষা করা হয় (Po No,
#     Generic Name, PCS IN CTN, Divider Measurement, Qty ইত্যাদি hidden থাকে)।
#
# ম্যাপিং (ইউজার-কনফার্মড):
#   Style Number          -> Gmt. Style No
#   Po No                 -> সবসময় 'N/A' (ওরা দেয়নি; hidden কলামের মান ব্যবহার হয় না)
#   Color Name            -> Reference/SKU Number
#   Size                  -> Pack Type
#   Carton Measuregement  -> Length/Width/Height ('L-46cm,W-33cm,H-22cm')
#   Carton Booking Qty    -> Order Qty
#   Item Name             -> UI সিলেকশন (Master Carton ইত্যাদি), Ply -> UI manual_ply
#
# 'Top & Bottom' রো (TOTAL CARTON QTY-এর নিচে):
#   Item Name = 'Top Bottom', Ply = 3, Style = ওই রো-র Style Number,
#   Length/Width = 'L- 43.5 cm,W- 30.5 cm' থেকে (Height থাকে না),
#   PO / Reference / Pack Type = 'N/A'।
#   একই মাপের একাধিক Top & Bottom রো থাকলে সেগুলো আলাদা লাইনই থাকে
#   (ফাইলে যেভাবে আছে সেভাবেই)।
# ---------------------------------------------------------------------------

NA = 'N/A'


def _norm(s):
    return re.sub(r'[^a-z0-9]', '', str(s or '').lower())


def _clean(v):
    if v is None:
        return ''
    return re.sub(r'\s+', ' ', str(v)).strip()


def _is_num(v):
    if v is None or isinstance(v, bool):
        return False
    try:
        float(v)
        return True
    except (TypeError, ValueError):
        return False


def _fmt_num(v):
    if v is None or v == '':
        return ''
    try:
        f = float(v)
    except (TypeError, ValueError):
        return ''
    return str(int(f)) if f == int(f) else str(round(f, 3))


def _hidden_col_indexes(ws):
    hidden = set()
    for dim in ws.column_dimensions.values():
        if dim.hidden:
            lo = dim.min or 0
            hi = dim.max or lo
            for idx in range(lo, hi + 1):
                hidden.add(idx)
    return hidden


def _hidden_row_indexes(ws):
    return {r for r, dim in ws.row_dimensions.items() if dim.hidden}


# 'L-46cm,W-33cm,H-22cm' / 'L- 43.5 cm,W- 30.5 cm' (Height ঐচ্ছিক)
_L_RE = re.compile(r'L\s*[-:]?\s*(\d+\.?\d*)', re.I)
_W_RE = re.compile(r'W\s*[-:]?\s*(\d+\.?\d*)', re.I)
_H_RE = re.compile(r'H\s*[-:]?\s*(\d+\.?\d*)', re.I)


def _parse_measurement(text):
    if not text:
        return '', '', ''
    t = str(text)
    l = _L_RE.search(t)
    w = _W_RE.search(t)
    h = _H_RE.search(t)
    return (
        _fmt_num(l.group(1)) if l else '',
        _fmt_num(w.group(1)) if w else '',
        _fmt_num(h.group(1)) if h else '',
    )


def _find_header(ws, hidden_cols, max_scan=30):
    """('stylenumber' + 'cartonbookingqty') দুটোই আছে এমন রো খুঁজে বের করে,
    এবং হেডিং -> কলাম-ইনডেক্স ম্যাপ রিটার্ন করে (hidden কলাম বাদ দিয়ে)।"""
    for r in range(1, min(ws.max_row, max_scan) + 1):
        heads = {}
        for c in range(1, ws.max_column + 1):
            if c in hidden_cols:
                continue
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            heads.setdefault(_norm(v), c)
        if 'stylenumber' in heads and 'cartonbookingqty' in heads:
            return r, heads
    return None, None


def _pick(heads, *prefixes):
    """normalize করা হেডিং-এর যেকোনো একটা prefix মিললে সেই কলাম রিটার্ন।
    (যেমন 'cartonmeas' দিয়ে 'Carton Measuregement' / 'Carton Measurement' দুটোই ধরা যায়)"""
    for key, col in heads.items():
        for p in prefixes:
            if key.startswith(p):
                return col
    return None


def read_southern_style_excel(file_stream, filename='', item_name_override='Master Carton', manual_ply=''):
    """মূল entry point। এই ফরম্যাট না হলে (কোনো শিটেই হেডার না পেলে) খালি
    লিস্ট [] রিটার্ন করে।"""
    wb = load_workbook(file_stream, data_only=True)
    all_items = []
    default_item_name = item_name_override or 'Master Carton'
    ply_value = manual_ply.strip() if manual_ply else NA

    for sn in wb.sheetnames:
        ws = wb[sn]
        if ws.sheet_state != 'visible':
            continue  # হাইড/ভেরি-হাইড শীট বাদ

        hidden_cols = _hidden_col_indexes(ws)
        hidden_rows = _hidden_row_indexes(ws)
        header_row, heads = _find_header(ws, hidden_cols)
        if header_row is None:
            continue  # এই শিট এই ফরম্যাটের না

        col_style = heads.get('stylenumber')
        col_color = _pick(heads, 'colorname', 'colour')
        col_size = _pick(heads, 'size')
        col_meas = _pick(heads, 'cartonmeas')
        col_qty = heads.get('cartonbookingqty')
        col_desc = _pick(heads, 'colorname', 'colour')  # 'Top & Bottom' লেবেল এই কলামেই বসে

        def val(r, c):
            return ws.cell(row=r, column=c).value if c else None

        last_style = ''
        for r in range(header_row + 1, ws.max_row + 1):
            if r in hidden_rows:
                continue

            label = _clean(val(r, col_desc))
            label_norm = _norm(label)
            style_cell = _clean(val(r, col_style))
            if style_cell:
                last_style = style_cell
            style_no = style_cell or last_style or NA

            # TOTAL CARTON QTY / TOTAL TOP & BOTTOM QTY সামারি রো — বাদ
            if label_norm.startswith('total'):
                continue

            qty_raw = val(r, col_qty)
            if not _is_num(qty_raw) or float(qty_raw) <= 0:
                continue
            qty = round(float(qty_raw))

            length, width, height = _parse_measurement(_clean(val(r, col_meas)))
            if not length:
                continue

            is_top_bottom = 'top' in label_norm and 'bottom' in label_norm

            if is_top_bottom:
                all_items.append({
                    'item_name': 'Top Bottom',
                    'ewo_no': NA,
                    'style_no': style_no,
                    'po_no': NA,
                    'length': length,
                    'width': width,
                    'height': '',
                    'ply': '3',
                    'qty': qty,
                    'pack_type': NA,
                    'reference': NA,
                    'color': '',
                    'size': '',
                    'delivery_date': '',
                    'measurement_unit': 'Cm',
                    'delivery_place_pdf': '',
                    'delivery_address_pdf': '',
                    '_sheet': sn,
                    '_source_file': filename,
                })
            else:
                # Master Carton — Style অবশ্যই থাকতে হবে
                if not style_cell:
                    continue
                all_items.append({
                    'item_name': default_item_name,
                    'ewo_no': NA,
                    'style_no': style_no,
                    'po_no': NA,
                    'length': length,
                    'width': width,
                    'height': height,
                    'ply': ply_value,
                    'qty': qty,
                    'pack_type': _clean(val(r, col_size)),
                    'reference': label,
                    'color': '',
                    'size': '',
                    'delivery_date': '',
                    'measurement_unit': 'Cm',
                    'delivery_place_pdf': '',
                    'delivery_address_pdf': '',
                    '_sheet': sn,
                    '_source_file': filename,
                })

    return all_items
