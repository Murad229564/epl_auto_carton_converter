import re
import pandas as pd

# ---------------------------------------------------------------------------
# Ventura (HK) Trading Limited — একই '排箱單 (PACKING LIST)' টেমপ্লেট-ফ্যামিলি,
# কিন্তু বায়ার-ভেদে Style/Reference/Pack Type কোন কলাম থেকে আসবে তা ভিন্ন।
# তাই buyer_key প্যারামিটার (UI থেকে ইউজার যে buyer সিলেক্ট করেছেন) দিয়ে
# ঠিক করা হয় কোন কলাম-লেবেল ব্যবহার হবে — কিন্তু কলামের actual পজিশন সবসময়
# header-লেবেল স্ক্যান করে ডাইনামিকভাবে বের করা হয় (fixed index ধরা হয় না)।
#
#   - প্রতিটা শিটে (এমনকি একই শিটের ভেতরেও, যেমন Kate Spade) একাধিক
#     "PURCHASE ORDER" ব্লক থাকতে পারে — প্রতিটা ব্লকের নিজস্ব হেডার-রো
#     (Carton No./TTL Ctns/... ইত্যাদি) থাকে, তাই পুরো শিট স্ক্যান করে
#     প্রতিটা ব্লক আলাদাভাবে ধরা হয়।
#   - PO নম্বর কখনো একই সেলে লেবেল+ভ্যালু একসাথে (Kate Spade: "PURCHASE
#     ORDER : 4520544711"), কখনো লেবেল আর ভ্যালু আলাদা সেলে (COACH/Michael
#     Kors/Le Sportsac/Vera Bradley) — দুটো কেসই হ্যান্ডেল করা হয়েছে।
#   - Measurement কখনো CM-এ, কখনো MM-এ থাকে — MM হলে CM-এ কনভার্ট করা হয়
#     (÷10)। Vera Bradley-র মতো কিছু শিটে CM আর INCH দুটো measurement-গ্রুপ
#     পাশাপাশি থাকে — তখন সবসময় CM-গ্রুপটাই নেওয়া হয়, INCH উপেক্ষা করা হয়।
#   - Qty সবসময় "TTL Ctns (CTNS)" কলাম থেকে।
#   - হাইড করা শিট আগে থেকেই স্কিপ করা হতো। এখন একই শিটের ভেতরে হাইড করা
#     রো বা কলামও (ইউজার-কনফার্মড) সম্পূর্ণ বাদ দেওয়া হয় — শুধু visible
#     রো/কলাম থেকেই ডাটা নেওয়া হবে। .xlsx-এ openpyxl দিয়ে, .xls-এ xlrd-এর
#     formatting_info দিয়ে হাইড রো/কলাম শনাক্ত করা হয়; শনাক্ত করতে না পারলে
#     (এক্সসেপশন হলে) নিরাপদ ডিফল্ট হিসেবে কিছুই হাইড ধরা হয় না।
# ---------------------------------------------------------------------------


def _norm(v):
    s = re.sub(r'[^a-z0-9]', '', str(v or '').lower())
    return s.replace('colour', 'color')  # British/American বানান একীভূত করা


def _clean(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ''
    return re.sub(r'\s+', ' ', str(v)).strip()


def _is_num(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return False
    return isinstance(v, (int, float))


# বায়ার-ভেদে reference/pack_type/style_no কোন লেবেল থেকে আসবে তার প্রোফাইল।
# style_no-এর জন্য চেষ্টার ক্রম: 'fty_hash' (শুধু Kate Spade-এ থাকে) না পেলে
# 'fty_style_combined' (বাকি সবগুলোতে থাকে)।
BUYER_PROFILES = {
    'katespade': {'ref': ['nrf_hash'], 'pack': ['color']},
    'michaelkors': {'ref': ['colorcode'], 'pack': ['colorname']},
    'coach': {'ref': ['color'], 'pack': ['colorname']},
    'lesportsac': {'ref': ['descriptionstylename', 'description'], 'pack': ['colorname']},
    'verabradley': {'ref': ['colorcode'], 'pack': ['colorname']},
}


def _find_header_rows(df, hidden_rows=None):
    """পুরো শিট স্ক্যান করে সব 'Carton No. + TTL Ctns' হেডার-রো (একাধিক
    PO-ব্লক থাকলে একাধিক পাওয়া যাবে) বের করে। হাইড করা রো-কে হেডার হিসেবে
    ধরা হয় না।"""
    hidden_rows = hidden_rows or set()
    rows = []
    n_rows, n_cols = df.shape
    for r in range(n_rows):
        if r in hidden_rows:
            continue
        labels = {_norm(df.iat[r, c]) for c in range(n_cols) if _clean(df.iat[r, c])}
        if 'cartonno' in labels and 'ttlctns' in labels:
            rows.append(r)
    return rows


def _extract_po_for_block(df, header_row):
    """header_row-এর উপরে (max ১৫ রো) 'purchase order' লেবেল খুঁজে PO
    নম্বর বের করে — লেবেল+ভ্যালু একই সেলে থাকলে সেখান থেকেই, নাহলে একই
    রো-তে ডানপাশের প্রথম নন-এম্পটি সেল থেকে।"""
    n_cols = df.shape[1]
    for r in range(max(0, header_row - 15), header_row):
        for c in range(n_cols):
            v = _clean(df.iat[r, c])
            if v and 'purchase order' in v.lower():
                after_colon = v.split(':', 1)[1].strip() if ':' in v else ''
                if after_colon:
                    return after_colon
                for c2 in range(c + 1, n_cols):
                    v2 = _clean(df.iat[r, c2])
                    if v2:
                        return v2
    return ''


def _build_col_map(df, header_row, hidden_cols=None):
    """header_row আর header_row+1 (দুই-রো হেডার) একসাথে স্ক্যান করে সব
    পরিচিত লেবেল-কলাম ম্যাপ করে। হাইড করা কলাম পুরোপুরি বাদ দেওয়া হয় —
    অর্থাৎ সেই কলামে যে লেবেলই থাকুক না কেন, সেটাকে কোনো ডাটা-সোর্স হিসেবে
    ধরা হবে না। Measurement-এর ক্ষেত্রে একাধিক গ্রুপ (CM/MM/INCH) থাকতে
    পারে, তাই সবগুলো আলাদাভাবে রেকর্ড করা হয় (হাইড গ্রুপ বাদে)।"""
    hidden_cols = hidden_cols or set()
    n_cols = df.shape[1]
    col_map = {}
    measurement_groups = []  # [(start_col, unit), ...]

    for c in range(n_cols):
        if c in hidden_cols:
            continue
        top = _clean(df.iat[header_row, c])
        sub = _clean(df.iat[header_row + 1, c]) if header_row + 1 < df.shape[0] else ''
        top_norm = _norm(top)
        combined = _norm(top + sub)
        if not top_norm:
            continue
        if top.replace(' ', '').upper() == 'FTY#':
            col_map['fty_hash'] = c
        elif top_norm == 'fty' and _norm(sub) == 'style':
            col_map['fty_style_combined'] = c
        elif top.replace(' ', '').upper() == 'NRF#':
            col_map['nrf_hash'] = c
        elif top_norm == 'colorcode':
            col_map['colorcode'] = c
        elif top_norm == 'colorname':
            col_map['colorname'] = c
        elif top_norm == 'color':
            col_map['color'] = c
        elif combined == 'descriptionstylename':
            col_map['descriptionstylename'] = c
        elif top_norm == 'description':
            col_map['description'] = c
        elif top_norm == 'ttlctns':
            col_map['qty'] = c
        elif top_norm == 'measurement':
            unit = 'mm' if 'mm' in _norm(sub) else ('inch' if 'inch' in _norm(sub) else 'cm')
            measurement_groups.append((c, unit))

    return col_map, measurement_groups


def _pick_measurement_group(measurement_groups):
    """CM গ্রুপ থাকলে সেটাই প্রেফার করে (INCH কখনো না), না থাকলে MM
    (কনভার্সন-সহ)।"""
    for c, unit in measurement_groups:
        if unit == 'cm':
            return c, False
    for c, unit in measurement_groups:
        if unit == 'mm':
            return c, True
    return None, False


def _get_visible_sheet_names(file_stream, filename):
    """হাইড (hidden/veryHidden) করা শিট বাদ দিতে, দৃশ্যমান শিটের নামের
    লিস্ট বের করে। শনাক্ত করতে না পারলে (কোনো এক্সসেপশন হলে) None রিটার্ন
    করে — তখন নিরাপদ ডিফল্ট হিসেবে সব শিটই প্রসেস হবে (আগের আচরণ)।"""
    try:
        file_stream.seek(0)
        if filename.lower().endswith('.xls'):
            import xlrd
            book = xlrd.open_workbook(file_contents=file_stream.read())
            return [
                name for name in book.sheet_names()
                if book.sheet_by_name(name).visibility == 0
            ]
        else:
            from openpyxl import load_workbook as _load_wb
            wb = _load_wb(file_stream, read_only=True, data_only=True)
            names = [ws.title for ws in wb.worksheets if ws.sheet_state == 'visible']
            wb.close()
            return names
    except Exception:
        return None
    finally:
        file_stream.seek(0)


def _get_hidden_rows_cols(file_stream, filename, sheet_name):
    """একটা নির্দিষ্ট (visible) শিটের ভেতরে হাইড করা রো/কলামের index
    (0-indexed, pandas-এর সাথে সামঞ্জস্যপূর্ণ) বের করে।

    .xlsx: openpyxl দিয়ে row_dimensions/column_dimensions-এর 'hidden'
    ফ্ল্যাগ চেক করা হয় (read_only মোডে এই তথ্য নির্ভরযোগ্যভাবে পাওয়া যায়
    না, তাই normal মোডে লোড করা হয়)।
    .xls: xlrd-এ formatting_info=True দিয়ে rowinfo_map/colinfo_map থেকে।

    শনাক্ত করতে না পারলে (এক্সসেপশন, বা ফরম্যাট সাপোর্ট না থাকলে) খালি
    সেট রিটার্ন করে — নিরাপদ ডিফল্ট হিসেবে তখন কিছুই হাইড ধরা হবে না।
    """
    hidden_rows = set()
    hidden_cols = set()
    try:
        file_stream.seek(0)
        if filename.lower().endswith('.xls'):
            import xlrd
            book = xlrd.open_workbook(file_contents=file_stream.read(), formatting_info=True)
            if sheet_name not in book.sheet_names():
                return hidden_rows, hidden_cols
            sheet = book.sheet_by_name(sheet_name)
            for rowx, rowinfo in getattr(sheet, 'rowinfo_map', {}).items():
                if getattr(rowinfo, 'hidden', 0):
                    hidden_rows.add(rowx)
            for colx, colinfo in getattr(sheet, 'colinfo_map', {}).items():
                if getattr(colinfo, 'hidden', 0):
                    hidden_cols.add(colx)
        else:
            from openpyxl import load_workbook as _load_wb
            from openpyxl.utils import column_index_from_string
            wb = _load_wb(file_stream, read_only=False, data_only=True)
            if sheet_name not in wb.sheetnames:
                wb.close()
                return hidden_rows, hidden_cols
            ws = wb[sheet_name]
            for row_idx, dim in ws.row_dimensions.items():
                if dim.hidden:
                    hidden_rows.add(row_idx - 1)  # openpyxl 1-indexed -> pandas 0-indexed
            for col_letter, dim in ws.column_dimensions.items():
                if dim.hidden:
                    try:
                        hidden_cols.add(column_index_from_string(col_letter) - 1)
                    except ValueError:
                        continue
            wb.close()
    except Exception:
        return set(), set()
    finally:
        file_stream.seek(0)
    return hidden_rows, hidden_cols


def read_ventura_style_excel(file_stream, filename='', buyer_key='', item_name_override='', manual_ply=''):
    """মূল entry point। buyer_key যেমন 'Kate Spade'/'Michael Kors'/'Coach'/
    'Le Sportsac'/'Vera Bradley' (case-insensitive, স্পেস বাদ দিয়ে) — কোন
    কলাম থেকে reference/pack_type নেবে তা ঠিক করতে ব্যবহার হয়। item_name
    আর ply এক্সেলে থাকে না, তাই UI থেকে (existing outhouse_item_name/
    outhouse_ply কনভেনশন অনুযায়ী) override হিসেবে নেওয়া হয়। এই ফরম্যাট
    না হলে (কোনো header block না পেলে) খালি লিস্ট [] রিটার্ন করে।

    হাইড করা শিট, এবং একই শিটের ভেতরে হাইড করা রো/কলাম — দুটোই স্কিপ করা
    হয়, শুধু visible ডাটা থেকেই লাইন-আইটেম তৈরি হয়।"""
    profile_key = _norm(buyer_key)
    profile = BUYER_PROFILES.get(profile_key, {})
    ref_labels = profile.get('ref', [])
    pack_labels = profile.get('pack', [])

    sheets = pd.read_excel(file_stream, sheet_name=None, header=None)
    visible_sheet_names = _get_visible_sheet_names(file_stream, filename)
    all_items = []

    for sheet_name, df in sheets.items():
        if visible_sheet_names is not None and sheet_name not in visible_sheet_names:
            continue  # হাইড করা শিট — স্কিপ

        hidden_rows, hidden_cols = _get_hidden_rows_cols(file_stream, filename, sheet_name)

        header_rows = _find_header_rows(df, hidden_rows)
        for idx, header_row in enumerate(header_rows):
            col_map, measurement_groups = _build_col_map(df, header_row, hidden_cols)
            if 'qty' not in col_map:
                continue

            style_col = col_map.get('fty_hash', col_map.get('fty_style_combined'))
            if style_col is None:
                continue

            ref_col = next((col_map[k] for k in ref_labels if k in col_map), None)
            pack_col = next((col_map[k] for k in pack_labels if k in col_map), None)

            meas_col, needs_mm_conversion = _pick_measurement_group(measurement_groups)
            if meas_col is None:
                continue

            po_no = _extract_po_for_block(df, header_row)

            next_header_row = header_rows[idx + 1] if idx + 1 < len(header_rows) else df.shape[0]
            qty_col = col_map['qty']

            for r in range(header_row + 2, next_header_row):
                if r >= df.shape[0]:
                    break
                if r in hidden_rows:
                    continue  # হাইড করা রো — স্কিপ
                qty_val = df.iat[r, qty_col]
                style_val = _clean(df.iat[r, style_col])
                if not _is_num(qty_val) or not style_val:
                    continue

                def fmt(n):
                    return str(int(n)) if float(n).is_integer() else str(n)

                def parse_dim(col_offset):
                    """L/W/H যেকোনো একটা কলাম আলাদাভাবে পার্স করে — ফাঁকা/সংখ্যা
                    না হলে '' রিটার্ন করে, পুরো রো বাদ দেয় না (Divider-জাতীয়
                    আইটেমে Height না-ও থাকতে পারে, সেটাও visible-যা-আছে-তাই
                    হিসেবে রাখা হয়)। যে কলাম হাইড, তার ভ্যালুও '' ধরা হয়।"""
                    dim_col = meas_col + col_offset
                    if dim_col in hidden_cols:
                        return ''
                    try:
                        v = float(df.iat[r, dim_col])
                    except (TypeError, ValueError):
                        return ''
                    if needs_mm_conversion:
                        v = v / 10
                    return fmt(v)

                length = parse_dim(0)
                width = parse_dim(2)
                height = parse_dim(4)
                if not length and not width and not height:
                    continue  # measurement কলামেই কিছু নেই — এই রো সত্যিই ডাটা-রো না

                all_items.append({
                    'item_name': item_name_override or '',
                    'ewo_no': 'N/A',
                    'style_no': style_val,
                    'po_no': po_no,
                    'length': length,
                    'width': width,
                    'height': height,
                    'ply': manual_ply.strip() if manual_ply else '',
                    'qty': qty_val,
                    'pack_type': _clean(df.iat[r, pack_col]) if pack_col is not None else 'N/A',
                    'reference': _clean(df.iat[r, ref_col]) if ref_col is not None else 'N/A',
                    'remarks': '',
                    'color': '',
                    'size': '',
                    'delivery_date': '',
                    'measurement_unit': 'Cm',
                    'delivery_place_pdf': '',
                    'delivery_address_pdf': '',
                    '_sheet': sheet_name,
                })

    return all_items